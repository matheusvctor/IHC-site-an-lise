import io
import re
import base64
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color="374151", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
      <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
    <w:tcMar {nsdecls("w")}>
      <w:top w:w="{top}" w:type="dxa"/>
      <w:bottom w:w="{bottom}" w:type="dxa"/>
      <w:left w:w="{left}" w:type="dxa"/>
      <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_formatted_text(p, text_node):
    """Auxiliar para adicionar texto com formatação (bold, italic, code)."""
    if isinstance(text_node, str):
        p.add_run(text_node)
        return

    for child in text_node.children:
        if isinstance(child, str):
            p.add_run(child)
        elif child.name in ['strong', 'b']:
            run = p.add_run(child.get_text())
            run.bold = True
        elif child.name in ['em', 'i']:
            run = p.add_run(child.get_text())
            run.italic = True
        elif child.name == 'code':
            run = p.add_run(child.get_text())
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            add_formatted_text(p, child)

def create_docx():
    print("Iniciando geração do documento DOCX...")
    with open('RELATORIO_CONFORMIDADE_ISO_9241_17.html', 'r', encoding='utf-8') as f:
        html_content = f.read()

    soup = BeautifulSoup(html_content, 'html.parser')
    doc = Document()

    # 1. Configurações de Página (ABNT NBR 14724)
    section = doc.sections[0]
    section.top_margin = Cm(3.0)
    section.left_margin = Cm(3.0)
    section.bottom_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    # Estilo Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    # ==================== CAPA ====================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_after = Pt(70)
    r = p_inst.add_run("UNIVERSIDADE ESTADUAL DA PARAÍBA\nCAMPUS I - CAMPINA GRANDE\nCENTRO DE CIÊNCIAS E TECNOLOGIA\nDEPARTAMENTO DE COMPUTAÇÃO\nCURSO DE BACHARELADO EM CIÊNCIA DA COMPUTAÇÃO")
    r.bold = True
    r.font.size = Pt(11)

    p_autor = doc.add_paragraph()
    p_autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_autor.paragraph_format.space_after = Pt(110)
    r_autor = p_autor.add_run("MATHEUS VICTOR")
    r_autor.bold = True
    r_autor.font.size = Pt(12)

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.line_spacing = 1.2
    p_titulo.paragraph_format.space_after = Pt(12)
    r_tit = p_titulo.add_run("AVALIAÇÃO DE CONFORMIDADE ERGONÔMICA DE USABILIDADE DO SOFTWARE SYNAPSE HEALTH FRENTE À NORMA INTERNACIONAL ISO 9241-17")
    r_tit.bold = True
    r_tit.font.size = Pt(14)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.line_spacing = 1.2
    p_sub.paragraph_format.space_after = Pt(150)
    r_sub = p_sub.add_run("Diagnóstico Ergonômico, Identificação de Não Conformidades, Refatoração de Interfaces e Análise Comparativa de Aderência")
    r_sub.bold = True
    r_sub.font.size = Pt(11)

    p_cidade = doc.add_paragraph()
    p_cidade.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cidade.paragraph_format.line_spacing = 1.15
    r_cid = p_cidade.add_run("CAMPINA GRANDE - PB\n2026")
    r_cid.bold = True
    r_cid.font.size = Pt(11)

    doc.add_page_break()

    # ==================== FOLHA DE ROSTO ====================
    p_fr_autor = doc.add_paragraph()
    p_fr_autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fr_autor.paragraph_format.space_after = Pt(120)
    r_fr_aut = p_fr_autor.add_run("MATHEUS VICTOR")
    r_fr_aut.bold = True
    r_fr_aut.font.size = Pt(12)

    p_fr_tit = doc.add_paragraph()
    p_fr_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fr_tit.paragraph_format.line_spacing = 1.2
    p_fr_tit.paragraph_format.space_after = Pt(10)
    r_fr_t = p_fr_tit.add_run("AVALIAÇÃO DE CONFORMIDADE ERGONÔMICA DE USABILIDADE DO SOFTWARE SYNAPSE HEALTH FRENTE À NORMA INTERNACIONAL ISO 9241-17")
    r_fr_t.bold = True
    r_fr_t.font.size = Pt(13)

    p_fr_sub = doc.add_paragraph()
    p_fr_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fr_sub.paragraph_format.space_after = Pt(80)
    r_fr_s = p_fr_sub.add_run("Diagnóstico Ergonômico, Identificação de Não Conformidades, Refatoração de Interfaces e Análise Comparativa de Aderência")
    r_fr_s.bold = True
    r_fr_s.font.size = Pt(11)

    p_nat = doc.add_paragraph()
    p_nat.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_nat.paragraph_format.left_indent = Cm(8.0)
    p_nat.paragraph_format.line_spacing = 1.0
    p_nat.paragraph_format.space_after = Pt(130)
    r_nat = p_nat.add_run("Relatório técnico-científico apresentado ao Curso de Bacharelado em Ciência da Computação do Centro de Ciências e Tecnologia da Universidade Estadual da Paraíba (UEPB), como requisito parcial de avaliação na disciplina de Interface Homem-Computador (IHC).\n\nDocente: Prof. Dr. Daniel Scherer")
    r_nat.font.size = Pt(10)

    p_fr_cid = doc.add_paragraph()
    p_fr_cid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fr_cid.paragraph_format.line_spacing = 1.15
    r_fr_c = p_fr_cid.add_run("CAMPINA GRANDE - PB\n2026")
    r_fr_c.bold = True
    r_fr_c.font.size = Pt(11)

    doc.add_page_break()

    # ==================== RESUMO ====================
    h_res = doc.add_paragraph()
    h_res.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_res.paragraph_format.space_after = Pt(14)
    r_h_res = h_res.add_run("RESUMO")
    r_h_res.bold = True
    r_h_res.font.size = Pt(12)

    p_res = doc.add_paragraph()
    p_res.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_res.paragraph_format.line_spacing = 1.5
    p_res.paragraph_format.first_line_indent = Cm(0)
    p_res.paragraph_format.space_after = Pt(14)
    p_res.add_run("Este relatório apresenta a avaliação de conformidade de usabilidade do software ")
    r_bold1 = p_res.add_run("Synapse Health")
    r_bold1.bold = True
    p_res.add_run(", ecossistema clínico voltado ao gerenciamento de atendimentos, recepção e agendamento ambulatorial, com base na norma internacional ")
    r_bold2 = p_res.add_run("ISO 9241-17")
    r_bold2.bold = True
    p_res.add_run(", que estabelece requisitos e recomendações ergonômicas para diálogos de preenchimento de formulários em terminais com telas de apresentação visual. Foram analisadas minuciosamente as telas de ")
    p_res.add_run("Cadastro de Paciente, Cadastro de Médico Especialista e Agendamento & Triagem Clínica, ")
    p_res.add_run("adotando-se o procedimento de avaliação em duas etapas estabelecido no Anexo A da norma: determinação de aplicabilidade e avaliação de aderência. O trabalho descreve os módulos do sistema e a variabilidade de seus 24 campos de preenchimento distribuídos em 14 tipos distintos de controles HTML5, registra o catálogo de 24 não conformidades identificadas na versão preliminar (falhas estruturais, de entrada, feedback bloqueante e perda de dados por navegação) e documenta as intervenções técnicas de código realizadas para assegurar plena conformidade. Por meio de evidências visuais de alta fidelidade e quadros analíticos comparativos, comprova-se a evolução do sistema de uma Taxa de Aderência inicial de 46,7% para 100,0% na versão final corrigida.")

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.first_line_indent = Cm(0)
    r_kw_lbl = p_kw.add_run("Palavras-chave: ")
    r_kw_lbl.bold = True
    p_kw.add_run("Interação Humano-Computador. ISO 9241-17. Usabilidade. Preenchimento de Formulários. Synapse Health.")

    doc.add_page_break()

    # ==================== LISTA DE FIGURAS ====================
    h_lf = doc.add_paragraph()
    h_lf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_lf.paragraph_format.space_after = Pt(14)
    r_hlf = h_lf.add_run("LISTA DE FIGURAS")
    r_hlf.bold = True

    figs_list = [
        ("Figura 1", "Tela de Cadastro de Paciente: versão inicial preliminar (MedFlow Clinic OS)", "17"),
        ("Figura 2", "Tela de Cadastro de Paciente: versão final após adequações de conformidade", "18"),
        ("Figura 3", "Validação em tempo real e painel superior de pendências (Paciente)", "19"),
        ("Figura 4", "Mecanismo de cancelamento reversível com banner Desfazer Limpeza (Paciente)", "19"),
        ("Figura 5", "Tela de Cadastro de Médico Especialista: versão preliminar inicial", "20"),
        ("Figura 6", "Tela de Cadastro de Médico Especialista: versão final com unidade anos integrada", "21"),
        ("Figura 7", "Validação inline e bloqueio de CRM duplicado na tela de Médico", "22"),
        ("Figura 8", "Tela de Agendamento & Triagem Clínica: versão inicial", "23"),
        ("Figura 9", "Formato de atendimento Presencial com sala virtual bloqueada e desabilitada", "24"),
        ("Figura 10", "Formato Telemedicina com link da sala virtual reativado e validado", "24"),
        ("Figura 11", "Contador de caracteres em tempo real e resumo de falhas no Agendamento", "25"),
    ]
    for tag, desc, pag in figs_list:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.line_spacing = 1.15
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.first_line_indent = Cm(0)
        p_item.add_run(f"{tag} – {desc} ")
        r_dots = p_item.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . ")
        r_dots.font.color.rgb = RGBColor(156, 163, 175)
        p_item.add_run(f" {pag}")

    doc.add_page_break()

    # ==================== LISTA DE QUADROS ====================
    h_lq = doc.add_paragraph()
    h_lq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_lq.paragraph_format.space_after = Pt(14)
    r_hlq = h_lq.add_run("LISTA DE QUADROS")
    r_hlq.bold = True

    quadros_list = [
        ("Quadro 1", "Mapeamento de Campos do Cadastro de Paciente", "8"),
        ("Quadro 2", "Mapeamento de Campos do Cadastro de Médico Especialista", "8"),
        ("Quadro 3", "Mapeamento de Campos do Agendamento e Triagem Clínica", "9"),
        ("Quadro 4", "Matriz de Variabilidade de Controles de Entrada HTML5", "10"),
        ("Quadro 5", "Correlação entre Cláusulas da ISO 9241-17 e Componentes do Synapse Health", "12"),
        ("Quadro 6", "Síntese da Avaliação Preliminar por Tela", "14"),
        ("Quadro 7", "Síntese da Avaliação Preliminar por Cláusula da ISO 9241-17", "14"),
        ("Quadro 8", "Catálogo de Não Conformidades Ergonômicas da Versão Inicial", "15"),
        ("Quadro 9", "Comparativo Antes vs. Depois – Cadastro de Paciente", "17"),
        ("Quadro 10", "Comparativo Antes vs. Depois – Cadastro de Médico Especialista", "20"),
        ("Quadro 11", "Comparativo Antes vs. Depois – Agendamento e Triagem Clínica", "23"),
        ("Quadro 12", "Matriz de Conformidade Final da ISO 9241-17", "26"),
        ("Quadro 13", "Checklist Geral de Aplicabilidade e Aderência (Tabela A.1 do Anexo A)", "27"),
    ]
    for tag, desc, pag in quadros_list:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.line_spacing = 1.15
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.first_line_indent = Cm(0)
        p_item.add_run(f"{tag} – {desc} ")
        r_dots = p_item.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . ")
        r_dots.font.color.rgb = RGBColor(156, 163, 175)
        p_item.add_run(f" {pag}")

    doc.add_page_break()

    # ==================== SUMÁRIO ====================
    h_sum = doc.add_paragraph()
    h_sum.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_sum.paragraph_format.space_after = Pt(14)
    r_hsum = h_sum.add_run("SUMÁRIO")
    r_hsum.bold = True

    sum_items = [
        ("1 INTRODUÇÃO E OBJETIVOS", "7", True),
        ("  1.1 Contextualização", "7", False),
        ("  1.2 Objetivo Geral", "7", False),
        ("  1.3 Objetivos Específicos", "7", False),
        ("2 O SISTEMA SYNAPSE HEALTH", "8", True),
        ("  2.1 Visão Geral do Sistema", "8", False),
        ("  2.2 Descrição dos Módulos e Mapeamento de Campos", "8", False),
        ("  2.3 Matriz de Variabilidade de Controles", "10", False),
        ("3 PROCEDIMENTO METODOLÓGICO DA AVALIAÇÃO", "11", True),
        ("  3.1 A Norma ISO 9241-17 e o Processo em Duas Etapas", "11", False),
        ("  3.2 Métodos de Verificação e Critérios de Aceitação", "11", False),
        ("  3.3 Escopo Funcional das Cláusulas Normativas", "12", False),
        ("4 DIAGNÓSTICO PRELIMINAR DE CONFORMIDADE", "13", True),
        ("  4.1 Síntese Quantitativa dos Problemas Identificados", "13", False),
        ("  4.2 Catálogo Detalhado de Não Conformidades (P01 a P24)", "15", False),
        ("5 REFATORAÇÃO DE INTERFACE E ANÁLISE COMPARATIVA", "16", True),
        ("  5.1 Intervenções Ergonômicas e Comparativo Antes vs. Depois", "16", False),
        ("  5.2 Matriz de Conformidade Final Consolidada", "26", False),
        ("  5.3 Checklist Oficial de Aplicabilidade e Aderência (Tabela A.1)", "27", False),
        ("6 CONCLUSÃO E CONSIDERAÇÕES FINAIS", "29", True),
        ("REFERÊNCIAS", "30", True),
    ]
    for title, pag, is_main in sum_items:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.line_spacing = 1.15
        p_s.paragraph_format.space_after = Pt(3)
        p_s.paragraph_format.first_line_indent = Cm(0)
        r_t = p_s.add_run(f"{title} ")
        if is_main:
            r_t.bold = True
        r_dots = p_s.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . ")
        r_dots.font.color.rgb = RGBColor(156, 163, 175)
        r_p = p_s.add_run(f" {pag}")
        if is_main:
            r_p.bold = True

    doc.add_page_break()

    # ==================== PROCESSAMENTO DO CORPO DO DOCUMENTO ====================
    # Vamos percorrer os elementos principais do HTML
    # Selecionamos os elementos após o sumário
    body_elements = soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'div'])

    # Filtramos para pegar os blocos da estrutura principal
    in_main_content = False

    for elem in body_elements:
        # Detectar início do conteúdo textual
        if elem.name == 'h1' and '1 INTRODUÇÃO E OBJETIVOS' in elem.get_text():
            in_main_content = True

        if not in_main_content:
            continue

        # Evitar processamento duplo de tags filhas já tratadas
        if elem.parent and elem.parent.get('class') and any(c in elem.parent.get('class') for c in ['quadro-wrapper', 'figure-wrapper', 'figure-img-box', 'page-capa', 'page-folha-rosto', 'page-pre-textual']):
            continue

        # 1. H1 (Seção Primária)
        if elem.name == 'h1':
            title_text = elem.get_text().strip()
            if not title_text or 'UNIVERSIDADE' in title_text or 'AVALIAÇÃO DE' in title_text or title_text in ['RESUMO', 'LISTA DE FIGURAS', 'LISTA DE QUADROS', 'SUMÁRIO']:
                continue
            
            # Seção break: adiciona quebra de página antes se não for a seção 1
            if '1 INTRODUÇÃO' not in title_text:
                doc.add_page_break()

            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.first_line_indent = Cm(0)
            h.paragraph_format.line_spacing = 1.2
            r = h.add_run(title_text)
            r.bold = True
            r.font.size = Pt(12)

        # 2. H2 (Seção Secundária)
        elif elem.name == 'h2':
            title_text = elem.get_text().strip()
            if not title_text:
                continue
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.first_line_indent = Cm(0)
            h.paragraph_format.line_spacing = 1.2
            r = h.add_run(title_text)
            r.bold = True
            r.font.size = Pt(11)

        # 3. H3 (Seção Terciária)
        elif elem.name == 'h3':
            title_text = elem.get_text().strip()
            if not title_text:
                continue
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.first_line_indent = Cm(0)
            h.paragraph_format.line_spacing = 1.2
            r = h.add_run(title_text)
            r.bold = True
            r.font.size = Pt(10.5)

        # 4. Parágrafos comuns
        elif elem.name == 'p':
            # Ignorar se estiver dentro de quadro ou figura
            if elem.find_parent('div', class_=['quadro-wrapper', 'figure-wrapper', 'page-capa', 'page-folha-rosto', 'page-pre-textual']):
                continue
            text = elem.get_text().strip()
            if not text:
                continue
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)

            # Recuo de primeira linha
            is_no_indent = 'no-indent' in elem.get('class', []) or 'REFERÊNCIAS' in elem.find_previous(['h1', 'h2']).get_text()
            if is_no_indent:
                p.paragraph_format.first_line_indent = Cm(0)
            else:
                p.paragraph_format.first_line_indent = Cm(1.25)

            add_formatted_text(p, elem)

        # 5. Listas não ordenadas
        elif elem.name == 'ul':
            if elem.find_parent('div', class_=['quadro-wrapper', 'figure-wrapper', 'page-capa', 'page-folha-rosto', 'page-pre-textual']):
                continue
            for li in elem.find_all('li', recursive=False):
                p_li = doc.add_paragraph(style='List Bullet')
                p_li.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_li.paragraph_format.line_spacing = 1.25
                p_li.paragraph_format.space_after = Pt(3)
                p_li.paragraph_format.first_line_indent = Cm(0)
                add_formatted_text(p_li, li)

        # 6. Quadros (Tabelas ABNT)
        elif elem.name == 'div' and 'quadro-wrapper' in elem.get('class', []):
            header_div = elem.find('div', class_='quadro-header')
            table_tag = elem.find('table')
            fonte_div = elem.find('div', class_='quadro-fonte')

            if not table_tag:
                continue

            # Cabeçalho do Quadro
            if header_div:
                p_qh = doc.add_paragraph()
                p_qh.paragraph_format.space_before = Pt(12)
                p_qh.paragraph_format.space_after = Pt(3)
                p_qh.paragraph_format.first_line_indent = Cm(0)
                r_qh = p_qh.add_run(header_div.get_text().strip())
                r_qh.bold = True
                r_qh.font.size = Pt(9.5)

            # Processar Linhas da Tabela
            rows = table_tag.find_all('tr')
            if not rows:
                continue

            # Determinar número máximo de colunas
            max_cols = 0
            for r in rows:
                cols_count = len(r.find_all(['th', 'td']))
                if cols_count > max_cols:
                    max_cols = cols_count

            # Cria tabela Word
            doc_table = doc.add_table(rows=0, cols=max_cols)
            doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(doc_table, color="374151", sz="4", val="single")

            # Preencher linhas
            for r_idx, tr in enumerate(rows):
                row_cells_data = tr.find_all(['th', 'td'])
                # Se for Quadro 13 e for a primeira linha com colspan, normalizamos
                if len(row_cells_data) != max_cols and 'Quadro 13' in (header_div.get_text() if header_div else ''):
                    # Usamos cabeçalho padrão de 8 colunas para Quadro 13
                    if r_idx == 0:
                        doc_row = doc_table.add_row()
                        headers_q13 = ['Cláusula', 'Recomendação Resumida (ISO 9241-17)', 'Aplicabilidade', 'Método Apl.', 'Método Ader.', 'Antes', 'Depois', 'Evidência / Comentário no Synapse Health']
                        for c_idx, h_text in enumerate(headers_q13):
                            cell = doc_row.cells[c_idx]
                            set_cell_background(cell, "F3F4F6")
                            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                            p_c = cell.paragraphs[0]
                            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_c.paragraph_format.line_spacing = 1.0
                            p_c.paragraph_format.space_after = Pt(0)
                            r_c = p_c.add_run(h_text)
                            r_c.bold = True
                            r_c.font.size = Pt(8.0)
                    continue  # Pular segunda linha de cabeçalho do Quadro 13
                
                doc_row = doc_table.add_row()
                is_header_row = (r_idx == 0 and tr.find('th') is not None)

                for c_idx, cell_data in enumerate(row_cells_data):
                    if c_idx >= max_cols:
                        break
                    cell = doc_row.cells[c_idx]
                    set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                    
                    p_c = cell.paragraphs[0]
                    p_c.paragraph_format.line_spacing = 1.15
                    p_c.paragraph_format.space_after = Pt(0)
                    
                    cell_text = cell_data.get_text().strip()
                    # Alinhamento
                    if is_header_row or len(cell_text) <= 15:
                        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    if is_header_row or cell_data.name == 'th':
                        set_cell_background(cell, "F3F4F6")
                        r_c = p_c.add_run(cell_text)
                        r_c.bold = True
                        r_c.font.size = Pt(8.5)
                    else:
                        r_c = p_c.add_run(cell_text)
                        r_c.font.size = Pt(8.0)

            # Fonte do Quadro
            if fonte_div:
                p_qf = doc.add_paragraph()
                p_qf.paragraph_format.space_before = Pt(3)
                p_qf.paragraph_format.space_after = Pt(12)
                p_qf.paragraph_format.first_line_indent = Cm(0)
                r_qf = p_qf.add_run(fonte_div.get_text().strip())
                r_qf.font.size = Pt(8.5)
                r_qf.font.color.rgb = RGBColor(75, 85, 99)

        # 7. Figuras (Imagens ABNT)
        elif elem.name == 'div' and 'figure-wrapper' in elem.get('class', []):
            title_div = elem.find('div', class_='figure-title')
            img_tag = elem.find('img')
            fonte_div = elem.find('div', class_='figure-source')

            # Título da Figura
            if title_div:
                p_ft = doc.add_paragraph()
                p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ft.paragraph_format.space_before = Pt(14)
                p_ft.paragraph_format.space_after = Pt(4)
                p_ft.paragraph_format.first_line_indent = Cm(0)
                r_ft = p_ft.add_run(title_div.get_text().strip())
                r_ft.bold = True
                r_ft.font.size = Pt(9.5)

            # Imagem
            if img_tag and img_tag.get('src'):
                src = img_tag['src']
                img_stream = None
                if src.startswith('data:image'):
                    b64_data = src.split(',', 1)[1]
                    img_stream = io.BytesIO(base64.b64decode(b64_data))
                
                if img_stream:
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(2)
                    p_img.paragraph_format.space_after = Pt(2)
                    run_img = p_img.add_run()
                    run_img.add_picture(img_stream, width=Inches(5.7))

            # Fonte da Figura
            if fonte_div:
                p_ff = doc.add_paragraph()
                p_ff.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_ff.paragraph_format.space_before = Pt(3)
                p_ff.paragraph_format.space_after = Pt(14)
                p_ff.paragraph_format.first_line_indent = Cm(0)
                r_ff = p_ff.add_run(fonte_div.get_text().strip())
                r_ff.font.size = Pt(8.5)
                r_ff.font.color.rgb = RGBColor(75, 85, 99)

    out_file = 'RELATORIO_CONFORMIDADE_ISO_9241_17.docx'
    doc.save(out_file)
    print(f"Documento DOCX gerado com sucesso: {out_file}")

if __name__ == '__main__':
    create_docx()
